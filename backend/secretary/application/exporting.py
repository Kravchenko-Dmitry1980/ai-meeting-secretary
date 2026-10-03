from __future__ import annotations

import io
import json


def export_meeting(meeting, segments, summary, format, *, assignment_snapshot=None,raw_summary=None,summary_context=None):
    transcript = "\n".join(f"[{s['id']}] {s['channel']}; {s['timing_precision']}; {s['start_ms'] if s['start_ms'] is not None else '?'}–{s['end_ms'] if s['end_ms'] is not None else '?'} ms\n{s['text']}" for s in segments)
    if format == "json":
        return json.dumps({"meeting": meeting, "segments": segments, "summary": summary,
            'raw_summary':raw_summary if raw_summary is not None else summary,'summary_context':summary_context,
            'task_assignments':assignment_snapshot.model_dump(mode='json') if assignment_snapshot is not None else None}, ensure_ascii=False, indent=2).encode(), "application/json"
    title_prefix, section_prefix = ("", "") if format == "txt" else ("# ", "## ")
    lines = [title_prefix + meeting['title'], "", section_prefix + "Исходная расшифровка", transcript]
    if summary:
        if summary.get("excluded_items_count", 0):
            lines.extend(["", f"Проверка итогов: исключено неподтверждённых пунктов из черновиков: {summary['excluded_items_count']}. "
                          "Это количество пунктов черновиков, а не уникальных задач. Проверьте полноту итогов по исходной расшифровке."])
        lines.extend(["", section_prefix + "Краткое содержание", summary["overview"], "", section_prefix + "Читаемая версия", summary["readable_transcript"]])
        for field, label in (("decisions", "Решения"), ("action_items", "Задачи"), ("open_questions", "Открытые вопросы")):
            lines.extend(["", section_prefix + label])
            for item in summary[field]:
                extra = ""
                if field == "action_items":
                    extra = f"; ответственный: {item['owner'] or 'Не определён'}; срок: {item['due_date'] or 'не указан'}"
                    extra += f"; основание: {item.get('assignment_basis','unknown')}; статус назначения: {item.get('assignment_status','needs_review')}; participant_id: {item.get('participant_id') or 'unknown'}"
                    extra += '; подтверждение назначения не доказывает принятие обязательства'
                lines.append(f"- {item['text']}{extra} [источники: {', '.join(item['source_segment_ids'])}]")
                lines.append('  Цитата: '+(item.get('evidence_quote') or 'не сохранена; требуется проверка'))
                source_map = {s['id']:s for s in segments}
                for ref in item['source_segment_ids']:
                    source = source_map.get(ref)
                    timing = 'таймкод неизвестен'
                    if source and source.get('timing_precision') != 'unknown' and source.get('start_ms') is not None and source.get('end_ms') is not None:
                        timing = f"{source['start_ms']}–{source['end_ms']} ms; точность: {source['timing_precision']}"
                    lines.append(f'  Источник {ref}: {timing}')
                if item.get('assignment_reason_codes'):
                    lines.append('  Причины проверки: '+', '.join(item['assignment_reason_codes']))
                if item.get('assignment_provenance'):
                    previous = item['assignment_provenance']
                    lines.append(f"  Предложение переноса: summary {previous['summary_version']}, action {previous['action_id']}, revision {previous['assignment_revision']}; требуется новая ручная проверка")
        if assignment_snapshot is not None:
            lines.extend(['',f'Происхождение: transcript {assignment_snapshot.transcript_version}; summary {assignment_snapshot.summary_version}; assignment revision {assignment_snapshot.revision}; attribution {assignment_snapshot.attribution_revision}; roster {assignment_snapshot.roster_revision}',
                          f'Captured context: {assignment_snapshot.captured_context_hash or "legacy/unknown"}'])
    else:
        lines.extend(["", "Итоги ещё не готовы. Экспорт содержит доступную исходную расшифровку."])
    if format == "docx":
        from docx import Document
        document = Document()
        for line in lines:
            if line.startswith("# "):
                document.add_heading(line[2:], 0)
            elif line.startswith("## "):
                document.add_heading(line[3:], 1)
            else:
                document.add_paragraph(line)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    text = "\n".join(lines)
    return text.encode("utf-8"), "text/plain; charset=utf-8" if format == "txt" else "text/markdown; charset=utf-8"

